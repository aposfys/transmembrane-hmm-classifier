import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import matplotlib.pyplot as plt
import networkx as nx
import numpy as np

# Create output folder if it doesn't exist
os.makedirs('output', exist_ok=True)

# Initialize the Word document
doc = Document()

# Define some formatting styles
def add_heading(text, level=1):
    doc.add_heading(text, level=level)

def add_paragraph(text):
    doc.add_paragraph(text)

def add_captioned_figure(image_path, caption):
    doc.add_picture(image_path, width=Inches(5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph(f"Figure: {caption}", style='Caption')

# Title and Introduction
add_heading("Building and Evaluating a Profile HMM for Single-Pass Type I Transmembrane Proteins", level=0)

add_heading("Introduction", level=1)
add_paragraph("Hidden Markov Models (HMMs) are probabilistic models commonly used to capture sequential data structures, such as biological sequences (Durbin et al., 1998). In this project, we aim to detect single-pass type I transmembrane proteins, essential components of cellular membranes, using profile HMMs built from known protein sequences.")

# How HMMs work
add_heading("Theory: How HMMs Work", level=1)
add_paragraph("An HMM consists of states, transitions between states, and emissions (observed symbols). Profile HMMs extend this concept to model sequence families, where match, insert, and delete states represent conserved and variable regions across a family of sequences.")

# Create HMM diagram
def create_hmm_diagram():
    G = nx.DiGraph()
    G.add_edges_from([
        ("Start", "Match 1"), ("Match 1", "Match 2"), ("Match 2", "Match 3"),
        ("Match 1", "Insert 1"), ("Insert 1", "Match 2"),
        ("Match 2", "Insert 2"), ("Insert 2", "Match 3"),
        ("Match 3", "End")
    ])
    pos = nx.spring_layout(G, seed=42)
    plt.figure(figsize=(8,6))
    nx.draw(G, pos, with_labels=True, node_size=2000, node_color='lightblue', arrowsize=20)
    plt.title("Simplified Profile HMM Structure")
    plt.savefig('output/hmm_diagram.png')
    plt.close()

create_hmm_diagram()
add_captioned_figure('output/hmm_diagram.png', "Simplified HMM showing match and insert states.")

# Workflow diagram
def create_workflow_diagram():
    G = nx.DiGraph()
    steps = ["Download Proteins", "Cluster Sequences", "Align Sequences", "Build HMM", "Test and Evaluate"]
    G.add_edges_from(zip(steps, steps[1:]))
    pos = nx.spring_layout(G, seed=42)
    plt.figure(figsize=(8,5))
    nx.draw(G, pos, with_labels=True, node_size=2500, node_color='lightgreen', arrowsize=20)
    plt.title("Workflow of the HMM Project")
    plt.savefig('output/workflow.png')
    plt.close()

create_workflow_diagram()
add_captioned_figure('output/workflow.png', "Workflow steps followed during the project.")

# Results
add_heading("Evaluation Results", level=1)

# First model results
doc.add_paragraph("First model (full-length sequences) evaluation:")
table1 = doc.add_table(rows=1, cols=2)
row = table1.rows[0].cells
row[0].text = 'Metric'
row[1].text = 'Value'
metrics_first = {"Sensitivity": "1.00", "Specificity": "0.48", "True Positives": "9", "False Negatives": "0", "True Negatives": "58", "False Positives": "63"}
for k, v in metrics_first.items():
    row = table1.add_row().cells
    row[0].text = k
    row[1].text = v

doc.add_paragraph("\nImproved model (transmembrane region focus) evaluation:")
table2 = doc.add_table(rows=1, cols=2)
row = table2.rows[0].cells
row[0].text = 'Metric'
row[1].text = 'Value'
metrics_second = {"Sensitivity": "0.778", "Specificity": "0.868", "True Positives": "7", "False Negatives": "2", "True Negatives": "59", "False Positives": "9"}
for k, v in metrics_second.items():
    row = table2.add_row().cells
    row[0].text = k
    row[1].text = v

# Comparison bar chart
def create_comparison_chart():
    labels = ['Sensitivity', 'Specificity']
    first = [1.0, 0.48]
    second = [0.778, 0.868]
    x = np.arange(len(labels))
    width = 0.35

    plt.figure(figsize=(8,5))
    plt.bar(x - width/2, first, width, label='First Model')
    plt.bar(x + width/2, second, width, label='Improved Model')
    plt.ylabel('Score')
    plt.title('Model Performance Comparison')
    plt.xticks(x, labels)
    plt.ylim(0, 1.2)
    plt.legend()
    plt.grid(True)
    plt.savefig('output/comparison_chart.png')
    plt.close()

create_comparison_chart()
add_captioned_figure('output/comparison_chart.png', "Comparison of Sensitivity and Specificity between models.")

# Confusion Matrices
def create_confusion_matrix(tp, fn, fp, tn, filename):
    matrix = np.array([[tp, fn], [fp, tn]])
    fig, ax = plt.subplots()
    im = ax.imshow(matrix, cmap='Blues')

    for i in range(2):
        for j in range(2):
            ax.text(j, i, matrix[i, j], ha='center', va='center', color='black')

    ax.set_xticks(np.arange(2))
    ax.set_yticks(np.arange(2))
    ax.set_xticklabels(['Positive', 'Negative'])
    ax.set_yticklabels(['Positive', 'Negative'])
    plt.xlabel('Predicted')
    plt.ylabel('Actual')
    plt.title('Confusion Matrix')
    plt.savefig(filename)
    plt.close()

create_confusion_matrix(9, 0, 63, 58, 'output/conf_matrix_first.png')
create_confusion_matrix(7, 2, 9, 59, 'output/conf_matrix_second.png')

add_captioned_figure('output/conf_matrix_first.png', "Confusion Matrix for First Model.")
add_captioned_figure('output/conf_matrix_second.png', "Confusion Matrix for Improved Model.")

# Conclusions
add_heading("Conclusions", level=1)
add_paragraph("Initially, the model demonstrated excellent sensitivity but lower specificity, indicating a high number of false positives. Focusing on the conserved transmembrane regions led to a better balance between sensitivity and specificity.")

# References
add_heading("References", level=1)
add_paragraph("Durbin R, Eddy SR, Krogh A, Mitchison G. Biological Sequence Analysis: Probabilistic Models of Proteins and Nucleic Acids. Cambridge University Press; 1998.")
add_paragraph("Eddy SR. Profile hidden Markov models. Bioinformatics. 1998;14(9):755-763.")

# Save document
doc.save('output/final_hmm_report.docx')

print("\n✅ Final report generated at 'output/final_hmm_report.docx'.")
